import openpyxl
from pathlib import Path

file_path = Path("ПУ Кавказ.xlsx")

wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb[wb.sheetnames[0]]

print(f"Файл: {file_path.name}")
print(f"Лист: {wb.sheetnames[0]}")
print(f"Размер: {ws.max_row} строк x {ws.max_column} колонок\n")

print("=" * 80)
print("СОДЕРЖИМОЕ ФАЙЛА:")
print("=" * 80)

for row in range(1, min(15, ws.max_row + 1)):
    row_data = []
    for col in range(1, min(15, ws.max_column + 1)):
        val = ws.cell(row, col).value
        if val is None:
            row_data.append("---")
        elif isinstance(val, (int, float)):
            row_data.append(str(val))
        else:
            row_data.append(str(val)[:20])
    print(f"Строка {row:2d}: {' | '.join(row_data)}")

print("\n" + "=" * 80)
print("АНАЛИЗ ЗАГОЛОВКОВ:")
print("=" * 80)

# Ищем строку с "Итого" и "Остаток"
for row in range(1, min(11, ws.max_row + 1)):
    row_text = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row, col).value
        if val:
            row_text.append(str(val).strip())
    
    row_str = " | ".join(row_text)
    if "итого" in row_str.lower() or "остаток" in row_str.lower():
        print(f"\nСтрока {row} (возможный заголовок):")
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row, col).value
            if val:
                print(f"  Колонка {col}: '{val}'")

print("\n" + "=" * 80)
print("ПОИСК КОЛОНОК 'ИТОГО' И 'ОСТАТОК':")
print("=" * 80)

# Предполагаем что заголовки на строке 3-5
for header_row in range(3, 6):
    itogo_cols = []
    ostatok_col = None
    
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(header_row, col).value
        if cell_value:
            cell_str = str(cell_value).strip().lower()
            if "итого" in cell_str or "итог" in cell_str:
                itogo_cols.append(col)
            elif "остаток" in cell_str or "остат" in cell_str:
                ostatok_col = col
    
    if itogo_cols or ostatok_col:
        print(f"\nСтрока {header_row}:")
        print(f"  Колонки 'Итого': {itogo_cols}")
        print(f"  Колонка 'Остаток': {ostatok_col}")

wb.close()

